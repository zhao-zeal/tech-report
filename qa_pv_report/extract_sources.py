from __future__ import annotations

import ast
import hashlib
import json
from pathlib import Path

import fitz
import pandas as pd
from docx import Document
from pptx import Presentation


ROOT = Path(r"E:\电量预测赛道")
WORK = ROOT / "技术报告"
OUT = WORK / "qa_pv_report"
OUT.mkdir(exist_ok=True)


def extract_embedded_sources() -> dict[str, str]:
    source_path = WORK / "submission/utils/hyx_pv_pipeline.py"
    tree = ast.parse(source_path.read_text(encoding="utf-8"))
    found: dict[str, str] = {}
    for node in tree.body:
        if isinstance(node, ast.Assign) and len(node.targets) == 1:
            target = node.targets[0]
            if isinstance(target, ast.Name) and target.id in {"_SOURCE_V36", "_SOURCE_V48"}:
                value = ast.literal_eval(node.value)
                found[target.id] = value
                (OUT / f"{target.id[1:].lower()}.py").write_text(value, encoding="utf-8")
    if "_SOURCE_V48" in found:
        v48_tree = ast.parse(found["_SOURCE_V48"])
        for node in v48_tree.body:
            if isinstance(node, ast.Assign) and len(node.targets) == 1:
                target = node.targets[0]
                if isinstance(target, ast.Name) and target.id == "_RUNTIME_SOURCE":
                    runtime = ast.literal_eval(node.value)
                    found["_RUNTIME_SOURCE"] = runtime
                    (OUT / "source_v48_runtime.py").write_text(runtime, encoding="utf-8")
                    break
    return found


def pdf_text(path: Path) -> str:
    doc = fitz.open(path)
    blocks = []
    for i, page in enumerate(doc, start=1):
        blocks.append(f"\n===== PAGE {i} =====\n")
        blocks.append(page.get_text("text"))
    return "".join(blocks)


def docx_text(path: Path) -> str:
    doc = Document(path)
    out = []
    for p in doc.paragraphs:
        if p.text.strip():
            out.append(p.text)
    for ti, table in enumerate(doc.tables, start=1):
        out.append(f"\n[TABLE {ti}]")
        for row in table.rows:
            out.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(out)


def pptx_text(path: Path) -> str:
    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        out.append(f"\n===== SLIDE {i} =====")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                out.append(shape.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    out.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(out)


def data_audit() -> dict:
    data = WORK / "submission/datasets/pv_load_forecasting"
    load_path = data / "train_data/load_data/pv_load_train.csv"
    train_wx_path = data / "train_data/weather_data/pv_weather_train_hourly.csv"
    test_wx_path = data / "test_data/weather_data/pv_weather_test_hourly.csv"
    load = pd.read_csv(load_path)
    train_wx = pd.read_csv(train_wx_path)
    test_wx = pd.read_csv(test_wx_path)
    result = {
        "load": {
            "shape": list(load.shape),
            "columns": list(load.columns),
            "head": load.head(3).to_dict("records"),
            "pv_count": int(load["pv_id"].nunique()),
            "ta_count": int(load["ta_id"].nunique()),
            "date_min": str(load["date"].min()),
            "date_max": str(load["date"].max()),
            "duplicate_pv_date": int(load.duplicated(["pv_id", "date"]).sum()),
            "missing_by_column": load.isna().sum().astype(int).to_dict(),
            "negative_load": int((pd.to_numeric(load["load"], errors="coerce") < 0).sum()),
            "zero_load": int((pd.to_numeric(load["load"], errors="coerce") == 0).sum()),
            "pv_ta_pairs": int(load[["pv_id", "ta_id"]].drop_duplicates().shape[0]),
            "pv_with_multiple_ta": int((load.groupby("pv_id")["ta_id"].nunique() > 1).sum()),
            "mapping": load[["pv_id", "ta_id"]].drop_duplicates().sort_values("pv_id").to_dict("records"),
        },
        "train_weather": {
            "shape": list(train_wx.shape),
            "columns": list(train_wx.columns),
            "head": train_wx.head(3).to_dict("records"),
            "ta_count": int(train_wx["ta_id"].nunique()),
            "datetime_min": str(train_wx["datetime"].min()),
            "datetime_max": str(train_wx["datetime"].max()),
            "duplicate_ta_datetime": int(train_wx.duplicated(["ta_id", "datetime"]).sum()),
            "sentinel_9999": {c: int((train_wx[c] == 9999).sum()) for c in train_wx.columns if pd.api.types.is_numeric_dtype(train_wx[c])},
            "missing_by_column": train_wx.isna().sum().astype(int).to_dict(),
        },
        "test_weather": {
            "shape": list(test_wx.shape),
            "columns": list(test_wx.columns),
            "head": test_wx.head(3).to_dict("records"),
            "ta_count": int(test_wx["ta_id"].nunique()),
            "datetime_min": str(test_wx["valid_datetime"].min() if "valid_datetime" in test_wx else test_wx["datetime"].min()),
            "datetime_max": str(test_wx["valid_datetime"].max() if "valid_datetime" in test_wx else test_wx["datetime"].max()),
            "missing_by_column": test_wx.isna().sum().astype(int).to_dict(),
        },
    }
    return result


def main() -> None:
    embedded = extract_embedded_sources()
    sources = {
        "rules": ROOT / "电量预测赛道-赛题解读及评审规则说明0821.pdf",
        "platform_pdf": ROOT / "新型电力系统下电能量预测算法开发赛道 - baseline平台演示培训0824-0704.pdf",
        "data_doc": ROOT / "研发工具包/研发工具包/新型电力系统下电能量预测算法开发赛道 - 赛题数据说明文档.docx",
        "package_doc": ROOT / "研发工具包/研发工具包/新型电力系统下电能量预测算法开发赛道 - 模型封装打包指引.docx",
        "baseline_doc": ROOT / "研发工具包/研发工具包/新型电力系统下电能量预测算法开发赛道 - Baseline运行指南.docx",
        "pv_ppt": ROOT / "分布式光伏发电量预测.pptx",
    }
    for key, path in sources.items():
        if path.suffix.lower() == ".pdf":
            text = pdf_text(path)
        elif path.suffix.lower() == ".docx":
            text = docx_text(path)
        else:
            text = pptx_text(path)
        (OUT / f"{key}.txt").write_text(text, encoding="utf-8")
    audit = data_audit()
    (OUT / "data_audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    summary = {
        "embedded_sources": {k: len(v) for k, v in embedded.items()},
        "source_hashes": {k: hashlib.sha256(v.read_bytes()).hexdigest() for k, v in sources.items()},
        "data_audit": audit,
    }
    (OUT / "extraction_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
