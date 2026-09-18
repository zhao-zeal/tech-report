from __future__ import annotations

import io
import os
import re
import shutil
import token
import tokenize
import xml.etree.ElementTree as ET
from pathlib import Path

import fitz
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"E:\电量预测赛道\技术报告")
REFERENCE = ROOT / "output/docx/地市用电量预测初赛技术报告_20260916_代码排版修复版.docx"
SVG_SOURCE = ROOT / "qa_pv_report/pv_model_architecture.svg"
DOCX_OUT = ROOT / "output/docx/分布式光伏发电量预测初赛技术报告_20260918.docx"
SVG_OUT = ROOT / "output/figures/分布式光伏发电量预测模型结构图.svg"
PNG_OUT = ROOT / "output/figures/分布式光伏发电量预测模型结构图.png"

BLUE = "2F5597"
TEAL = "1D7074"
PALE_BLUE = "DCE6F1"
PALE_TEAL = "DDEBF7"
LIGHT = "F6F8FA"
GRID = "B7C9DD"
BLACK = RGBColor(0, 0, 0)


def set_run_font(run, latin="Times New Roman", east="宋体", size=12, bold=None, color=BLACK):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    run._element.rPr.rFonts.set(qn("w:ascii"), latin)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), latin)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="5"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        elem = borders.find(tag)
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), size)
        elem.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, keep=True):
    paragraph.paragraph_format.keep_with_next = keep
    paragraph.paragraph_format.widow_control = True


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11.5)
    normal.font.color.rgb = BLACK
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Pt(23)
    pf.line_spacing = 1.45
    pf.space_after = Pt(4)
    pf.widow_control = True

    title = doc.styles["Title"]
    title.font.name = "Times New Roman"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "方正小标宋简体")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(7)

    for name, size, east in (("Heading 1", 16, "黑体"), ("Heading 2", 14, "黑体"), ("Heading 3", 12.5, "黑体")):
        try:
            st = doc.styles[name]
        except KeyError:
            st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = normal
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), east)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = BLACK
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.widow_control = True
        st.paragraph_format.space_before = Pt(11 if name != "Heading 3" else 7)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.first_line_indent = Pt(0)
        st.paragraph_format.line_spacing = 1.15


def add_title(doc):
    p = doc.add_paragraph(style="Title")
    p.add_run("新型电力系统下电能量预测算法开发赛道")
    for text, size in (("初赛技术报告", 21), ("分布式光伏发电量预测", 24)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8 if size == 21 else 18)
        r = p.add_run(text)
        set_run_font(r, east="方正小标宋简体", size=size, bold=True)
    t = add_table(doc, ["项目", "内容"], [
        ["预测对象", "100个分布式光伏用户"],
        ["训练区间", "2024年10月1日—2025年9月30日"],
        ["预测区间", "2025年10月1日—10月31日"],
        ["特殊考核时段", "2025年10月4日—10月7日"],
        ["最终方案", "晴空包络与预报气象校正双分支融合模型"],
    ], widths=[1.35, 5.6])
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("2026年9月")
    set_run_font(r, east="宋体", size=12)


def add_heading(doc, text, level=1, page_break=False):
    if page_break:
        doc.add_page_break()
    p = doc.add_heading(text, level=level)
    set_keep_with_next(p)
    return p


def add_body(doc, text, bold_prefix=None, no_indent=False):
    p = doc.add_paragraph(style="Normal")
    if no_indent:
        p.paragraph_format.first_line_indent = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if bold_prefix and text.startswith(bold_prefix):
        a, b = text[:len(bold_prefix)], text[len(bold_prefix):]
        r = p.add_run(a)
        set_run_font(r, east="宋体", size=11.5, bold=True)
        r = p.add_run(b)
        set_run_font(r, east="宋体", size=11.5)
    else:
        r = p.add_run(text)
        set_run_font(r, east="宋体", size=11.5)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Pt(21)
    p.paragraph_format.hanging_indent = Pt(12)
    p.add_run("• ")
    r = p.add_run(text)
    set_run_font(r, east="宋体", size=11.2)
    return p


def add_equation(doc, text, number=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_together = True
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    mr = OxmlElement("m:r")
    mt = OxmlElement("m:t")
    mt.text = text
    mr.append(mt)
    omath.append(mr)
    omath_para.append(omath)
    p._p.append(omath_para)
    if number:
        r = p.add_run(f"    （{number}）")
        set_run_font(r, size=10.5)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for j, h in enumerate(headers):
        c = hdr.cells[j]
        set_cell_shading(c, PALE_BLUE)
        set_cell_margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            c.width = Inches(widths[j])
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(h))
        set_run_font(r, east="黑体", size=font_size, bold=True)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        if i % 2:
            for c in cells:
                set_cell_shading(c, "F8FAFC")
        for j, value in enumerate(row):
            c = cells[j]
            set_cell_margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                c.width = Inches(widths[j])
            p = c.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, east="宋体", size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_caption(doc, text, kind="图"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    set_keep_with_next(p, False)
    r = p.add_run(text)
    set_run_font(r, east="宋体", size=10.5, bold=False)
    return p


def add_table_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(4)
    set_keep_with_next(p)
    r = p.add_run(text)
    set_run_font(r, east="黑体", size=10.5, bold=False)


TOKEN_COLORS = {
    token.COMMENT: RGBColor(0x00, 0x80, 0x00),
    token.STRING: RGBColor(0xA3, 0x15, 0x15),
    token.NUMBER: RGBColor(0x09, 0x86, 0x58),
}


def colored_code_paragraph(cell, line):
    p = cell.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.keep_together = True
    if not line:
        p.add_run(" ")
        return
    try:
        toks = list(tokenize.generate_tokens(io.StringIO(line + "\n").readline))
    except tokenize.TokenError:
        toks = []
    col = 0
    for tok in toks:
        if tok.type in (token.ENDMARKER, token.NEWLINE, token.NL, token.ENCODING):
            continue
        if tok.start[0] != 1:
            continue
        start = tok.start[1]
        if start > col:
            r = p.add_run(line[col:start])
            set_run_font(r, latin="Consolas", east="微软雅黑", size=8.6)
        text = tok.string
        r = p.add_run(text)
        color = TOKEN_COLORS.get(tok.type, RGBColor(0, 0, 0))
        if tok.type == token.NAME and text in __import__("keyword").kwlist:
            color = RGBColor(0x00, 0x00, 0xFF)
        set_run_font(r, latin="Consolas", east="微软雅黑", size=8.6, color=color)
        col = tok.end[1]
    if col < len(line):
        r = p.add_run(line[col:])
        set_run_font(r, latin="Consolas", east="微软雅黑", size=8.6)


def add_code_block(doc, title, code):
    add_table_title(doc, title)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="D0D7DE", size="4")
    row = table.rows[0]
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)
    cell = row.cells[0]
    set_cell_shading(cell, "F6F8FA")
    set_cell_margins(cell, top=110, start=150, bottom=110, end=150)
    cell.width = Inches(6.9)
    cell._element.remove(cell.paragraphs[0]._element)
    for line in code.strip("\n").splitlines():
        colored_code_paragraph(cell, line.rstrip().replace("\t", "    "))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def prepare_figure():
    SVG_OUT.parent.mkdir(parents=True, exist_ok=True)
    tree = ET.parse(SVG_SOURCE)
    root = tree.getroot()
    styles = {
        "t": {"font-family": "Microsoft YaHei,Arial,sans-serif", "fill": "#172433"},
        "h": {"font-size": "34px", "font-weight": "700"},
        "m": {"font-size": "28px", "font-weight": "700"},
        "s": {"font-size": "23px", "fill": "#405266"},
        "box": {"fill": "#f8fbfd", "stroke": "#9fb4c8", "stroke-width": "2"},
        "blue": {"fill": "#edf4fa", "stroke": "#7da5c6", "stroke-width": "2.5"},
        "teal": {"fill": "#edf8f6", "stroke": "#6aa99f", "stroke-width": "2.5"},
        "mix": {"fill": "#e7f0f7", "stroke": "#557fa4", "stroke-width": "2.5"},
        "out": {"fill": "#f3f6f8", "stroke": "#8799a9", "stroke-width": "2.5"},
        "a": {"stroke": "#66839d", "stroke-width": "3", "fill": "none", "marker-end": "url(#arrow)"},
        "at": {"stroke": "#4d9187", "stroke-width": "3", "fill": "none", "marker-end": "url(#arrowTeal)"},
    }
    for elem in root.iter():
        classes = elem.attrib.pop("class", "").split()
        for cls in classes:
            for key, value in styles.get(cls, {}).items():
                elem.set(key, value)
    ET.register_namespace("", "http://www.w3.org/2000/svg")
    tree.write(SVG_OUT, encoding="utf-8", xml_declaration=True)
    data = SVG_OUT.read_bytes()
    svg_doc = fitz.open(stream=data, filetype="svg")
    page = svg_doc[0]
    pix = page.get_pixmap(matrix=fitz.Matrix(2.5, 2.5), alpha=False)
    pix.save(PNG_OUT)
    svg_doc.close()


def build_doc():
    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    prepare_figure()
    shutil.copy2(REFERENCE, DOCX_OUT)
    doc = Document(DOCX_OUT)
    clear_body(doc)
    configure_styles(doc)
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.78)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    sec.header_distance = Inches(0.3)
    sec.footer_distance = Inches(0.3)

    add_title(doc)

    add_heading(doc, "一、方案概述", 1, page_break=True)
    add_body(doc, "本任务面向100个分布式光伏用户，使用2024年10月1日至2025年9月30日的日发电量、用户装机容量、用户—气象区域映射和小时级气象数据，预测2025年10月1日至10月31日逐用户日发电量。提交结果覆盖100个用户、31个自然日，共3100个用户—日期键；2025年10月4日至10月7日为极端天气专项考核窗口。")
    add_body(doc, "最终方案命名为“晴空包络与预报气象校正双分支融合模型”。分支一以晴空辐射包络、日尺度辐射物理量、历史发电特征和用户级极端随机树为主体，并叠加两随机种子的时序卷积预测及低发电风险门控；分支二先把数值预报气象校正到历史实况气象表征，再用共享专家和用户级专家的教师—学生训练生成31天直接预测。顶层以固定配置0.75和0.25融合两个分支。")
    add_body(doc, "赛题评分按用户分别计算相对均方根误差型得分系数，再对100个用户取平均。月度考核满分20分，特殊时段考核满分10分，任务三合计满分30分。本文只陈述官方材料、最终训练与推理调用链、实际数据审计能够证明的内容。")
    add_equation(doc, "S₍u,W₎ = max{0, 1 − √[(1/n) Σₜ∈W ((yᵤ,ₜ − ŷᵤ,ₜ)/yᵤ,ₜ)²]}", "1")
    add_body(doc, "式中，u为光伏用户，W为月度或特殊时段评价集合，n为该集合中的样本数，y和ŷ分别为实际发电量与预测发电量。相应考核得分为该系数在用户维度的均值乘以20分或10分权重。")

    add_table_title(doc, "表1 任务范围与交付约束")
    add_table(doc, ["项目", "代码与材料核实结果"], [
        ["用户与区域", "pv_1—pv_100，共100个用户；ta_1—ta_9，共9个气象区域"],
        ["训练标签", "36500个用户—日期键；181个load缺失值、112个零值、无负值"],
        ["训练气象", "78840行＝9区域×365日×24小时；VIS中1个9999异常标识"],
        ["预测气象", "6696行＝9区域×31日×24小时，字段与训练气象一致"],
        ["提交契约", "pv_id、date、pred；date格式YYYY/MM/DD；3100行；键唯一、数值有限且非负"],
    ], widths=[1.25, 5.7])

    add_heading(doc, "二、技术方案", 1, page_break=True)
    add_heading(doc, "（一）数据处理", 2)
    add_body(doc, "训练标签文件包含pv_id、ta_id、date、load和installed_capacity。代码要求每个用户仅映射到一个气象区域、装机容量保持不变且为正；用户—日期键不得重复，训练期365个日期必须完整。实际数据中100个用户分别映射至9个ta_id，映射关系满足上述约束。")
    add_body(doc, "分支一对每个用户的load按1%和99%分位截断。缺失值与小于10⁻⁶的近零值先用同区域、同日期其他用户的有效均值修复；该均值无效时使用该用户历史中位数。此规则把无法用于对数目标训练的缺失或近零记录恢复为区域同步水平，同时保留后续预测的非负边界。分支二保留标签有效性掩码，并以用户正发电量中位数作为尺度参数。")
    add_body(doc, "小时气象读取同时兼容datetime和valid_datetime字段。代码建立9个区域、规定日期、24个小时的完整笛卡尔积，拒绝重复键、缺失小时和非有限值；数值9999按缺测标识转为NaN后再进入统计。训练气象读取截止2025年9月30日，预测气象只读取2025年10月。")
    add_body(doc, "两个分支独立返回pv_id、date、pred。顶层先分别按预期3100键重排，校验重复键、缺失键、意外键、无穷值、NaN和负值，再进行逐键融合。最终按pv_id数字后缀和日期排序，以UTF-8 BOM写入submit_result.csv。任何校验失败均抛出异常，不产生不完整提交。")

    add_heading(doc, "（二）小时气象聚合与物理特征构造", 2)
    add_body(doc, "原始气象变量包括温度TEM、相对湿度RHU、15分钟降水PRE_15m、累计降水PRE_acc、总辐射SR、总云量TCC、直接辐射SWDDIR、散射辐射SWDDIF、能见度VIS和风速WS。模型不把24小时序列简单压缩为单一均值，而是从日统计、日内时段、辐射组成和天气持续性四个层面构造特征。")
    add_table_title(doc, "表2 小时气象到日尺度的特征体系")
    add_table(doc, ["类别", "实际生效的特征"], [
        ["原始变量", "TEM、RHU、PRE_15m、PRE_acc、SR、TCC、SWDDIR、SWDDIF、VIS、WS"],
        ["聚合统计", "均值、最大/最小、标准差、极差、总量、峰值时刻、辐射高于日均的小时数"],
        ["分时段", "上午6—11时、午间10—14/15时、下午12—18时、夜间0—5时；均值或总量"],
        ["物理派生", "晴空日辐射、晴朗指数、直接/散射辐射比例、晴天时数、连续阴云时长"],
        ["交互特征", "辐射加权云量、高辐射温度、温湿交互、降水—云量、能见度—湿度、风—温度"],
        ["时间天文", "日序、月份、星期、正余弦周期、太阳赤纬、归一化地外辐射代理量"],
    ], widths=[1.22, 5.73])
    add_body(doc, "分支一还计算上午与下午辐射非对称性、峰值辐射时刻云量、日间降水、午间湿度和日间风速。晴天时数定义为小时晴朗指数大于0.7的小时数，连续阴云时长由小时晴朗指数小于0.3的最长连续区间表示。分支二把预报气象进一步组织为42维原始输入：9个日均、9个午间均值、9个日极差、9个日标准差、1个辐射异常标志和5个日历量。")

    add_heading(doc, "（三）晴空包络与天气状态建模", 2)
    add_body(doc, "分支一按ta_id、年内日序和小时建立晴空辐射包络。对每个区域的SR，在目标日序前后各10天的季节窗口内取90%分位数，并用首尾三年复制实现跨年连续窗口；缺口以前向和后向方式补齐。日尺度晴空上界则对各区域历史日均SR在前后12个样本窗口内取90%分位，并作7点平滑。两种包络分别服务于小时曲线形态与日尺度季节上界。")
    add_body(doc, "预测期的夜间辐射修复仅作用于曲线视图：当已拟合晴空小时包络低于5 W/m²时，将SR、SWDDIR和SWDDIF置零，再重新计算cs_daily、kt_daily、kt_noon、kt_morning、kt_afternoon、cloud_run和sunny_hours七个字段。其余85个树模型输入、无曲线视图和时序卷积分支不被修改。代码同时要求各晴朗指数不超过5，以阻断夜间伪辐射放大。")
    add_body(doc, "天气状态由晴朗指数、云量、降水、风速和辐射组成共同刻画。树模型输出还乘以固定经验气象校准系数，该系数随日间风速和降水增加而下调，并限制在0.65—1.05。该校准是代码中的经验修正，不等同于理论辐射传输模型。")
    add_equation(doc, "c_wx = clip[0.94 − 0.021(W_day − 4) − 0.041{ln(1 + P_day) − 0.5}, 0.65, 1.05]", "2")

    add_heading(doc, "（四）历史发电特征与用户尺度处理", 2)
    add_body(doc, "树模型输入含1、2、3、7、14、28日滞后值，3、5、7、14、28日滚动均值，7日与14日滚动标准差，同星期前4周均值、7日短期趋势和用户月度发电水平。滞后与短窗统计描述近期状态，7日及同星期统计描述周周期，28日窗口和月度水平描述月内及季节变化。31天推理中，这组历史特征由训练截止时保存的最近60日队列一次性生成，未把未来真实发电量引入预测期。")
    add_body(doc, "用户差异通过三条路径进入模型：第一，分支一为每个用户单独训练全年与近180日两套极端随机树；第二，时序卷积和分支二均用用户正发电量中位数归一化并恢复到实际尺度；第三，分支二同时使用用户编号、气象区域编号、共享专家和用户级专家。该组合同时保留跨用户共性与装机规模、局部响应差异。")

    add_heading(doc, "（五）分支一模型设计：晴空包络辐射特征与序列校正", 2, page_break=True)
    add_body(doc, "该分支由用户级树模型、双随机种子时序卷积模型和低发电风险门控组成。用户级树模型使用92个特征，目标为ln(load＋0.01×installed_capacity)。每个用户训练一套全年森林和一套最近180日森林，共200个ExtraTreesRegressor；每套500棵树，min_samples_split＝4，min_samples_leaf＝2，max_features＝sqrt，random_state＝42。")
    add_body(doc, "树模型反变换为exp(预测值)－0.01×装机容量，再乘气象校准系数并截断为非负。10月4—7日使用近180日森林，其他日期使用全年森林并乘0.965。曲线视图与无曲线视图在普通日期按0.5/0.5组合，并把普通日期用户总量校准到无曲线视图总量；特殊时段保持无曲线视图。")
    add_body(doc, "时序部分以最近62天、100个用户的归一化发电量、有效性标志和9个气象量为历史输入，以未来9个气象量和5个日历量为预测输入。网络使用膨胀率1和3的时序卷积，训练两个随机种子（42、2026），推理取均值。普通日期按0.75×树模型组合＋0.25×时序卷积组合；10月4—7日在这一层保持树模型组合。")
    add_body(doc, "最后以时序卷积结果为风险排序信号：对每个用户选出预测最低的4天，且该结果低于当前组合值时，输出乘0.70；其余记录乘1.02。此门控未对10月4—7日设置日历保护，因此特殊窗口中的满足条件记录仍会被校正。分支输出再次执行3100行、键唯一、数值有限和非负校验。")

    add_heading(doc, "（六）分支二模型设计：预报气象校正与共享—用户专家蒸馏", 2)
    add_body(doc, "该分支先学习“预报气象→实况日天气”的桥接关系。历史实况来自ta_id对应的站点气象文件；映射在每个发布时点之前的180天内建立，综合温度变化相关性（0.5）、湿度相对15日趋势的相关性（0.3）和对数降水相关性（0.2）选择最相近站点。实况数据按物理范围清洗，每日有效小时少于20时不形成训练目标。")
    add_body(doc, "气象桥接模型以42维预报统计量和9维区域独热编码为输入，用160棵极端随机树同时预测17维实况目标。17维目标由温度均值/最高/最低/极差，湿度均值/最低/最高，日降水总量，以及温度、湿度、降水在日间、午间、夜间的9个统计量组成；降水维度采用ln(1＋x)。模型用树间方差与训练残差共同构造不确定性，并以历史样本外误差对各区域、月份、目标维度的标准差作0.5—4倍校正。")
    add_body(doc, "样本外气象特征按月滚动生成：每个月的桥接模型只使用该月首日之前的数据，训练起点覆盖2024年12月至2025年9月。该发布时点约束使学生模型训练所见的校正气象与正式预测时的可用信息结构一致，不使用2025年10月实况选择模型。")
    add_body(doc, "专家输入由原始预报统计、17维校正均值、17维不确定性、12维历史发电上下文、用户与区域身份、预测步长组成，并对26个气象字段增加前差、后差和三点均值，最终形成169维输入。历史上下文包含3、7、14、28、62日均值与标准差以及10%和90%分位数。")
    add_body(doc, "教师模型在历史月份使用实况气象，分别训练共享ExtraTrees与100个用户级ExtraTrees并取均值；学生模型使用当时可获得的校正预报气象，每14天构造最长31天样本，训练目标为真实归一化发电量对数与教师预测对数各0.5的混合。最终学生包含300棵树的共享模型和100个300棵树的用户模型，推理时两者各占0.5，再乘用户尺度恢复实际发电量。HistGradientBoosting仅存在于未被最终fit_student调用的兼容函数中，未写入最终方案。")

    add_heading(doc, "（七）双分支融合与约束输出", 2)
    add_body(doc, "顶层封装先加载并核对meta.json中的管线版本、固定权重、训练截止日期和策略说明，再分别调用两个分支。两个结果均按相同的pv_id—date预期索引重排，保证用户、日期、行顺序、预测单位和非负边界一致。")
    add_equation(doc, "ŷ = 0.75 × ŷ¹ + 0.25 × ŷ²", "3")
    add_body(doc, "0.75/0.25是最终代码写死的交付配置；本地材料未提供该权重为全局最优的实验证据。融合后再次检查100个用户、31个日期、3100条记录，无重复键、缺失键、意外键、NaN、无穷值和负值，并把日期格式化为YYYY/MM/DD。")

    add_heading(doc, "（八）关键难点与解决措施", 2)
    add_table_title(doc, "表3 关键难点、代码实施与作用边界")
    add_table(doc, ["关键难点", "实际实施的解决措施"], [
        ["小时气象与日标签尺度不同", "构造日统计、分时段、峰值、波动、持续性和辐射组成特征"],
        ["预报与实况存在系统偏差", "按发布时点滚动训练17维气象桥接模型，并输出均值与不确定性"],
        ["辐射、云量、降水关系强非线性", "晴空包络、晴朗指数、比例及交互特征配合ExtraTrees建模"],
        ["100个用户尺度与响应不同", "用户级森林、共享/用户专家、用户中位数尺度和身份特征协同"],
        ["训练历史仅一年", "跨用户共享、物理包络、近180日与全年模型并用，避免依赖跨年标签"],
        ["31天内无未来真实发电量", "树分支冻结训练期历史队列；学生模型按预测步长直接生成31天结果"],
        ["10月4—7日极端天气", "近180日森林、气象敏感特征与气象桥接；门控按实际风险信号执行"],
        ["精度、鲁棒性和封装平衡", "CPU树模型与批量预测、清单哈希、环境变量路径和多层输出校验"],
    ], widths=[1.75, 5.2], font_size=9.3)

    add_heading(doc, "（九）模型结构图", 2, page_break=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(PNG_OUT), width=Inches(5.55))
    add_caption(doc, "图1 分布式光伏发电量双分支融合模型结构")
    add_body(doc, "结构图从输入、特征处理、互补分支到融合输出对应最终调用链。分支一的树模型与时序卷积侧重晴空包络、历史发电状态和局部风险；分支二侧重预报气象偏差校正、发布时点约束与共享—用户专家协同。")

    add_heading(doc, "三、实施成果", 1, page_break=False)
    add_heading(doc, "（一）开发成果", 2)
    add_body(doc, "交付代码形成统一的train_pv_model和predict_pv_model接口。训练入口依次准备分支一的用户级森林、两随机种子时序卷积模型，以及分支二的气象桥接、教师与学生专家；顶层写入组合清单。推理入口只接收规定的10月气象文件与已训练模型，输出固定三列提交文件。")
    add_table_title(doc, "表4 模型文件、清单与目录组织")
    add_table(doc, ["层级", "文件或变量", "作用"], [
        ["顶层组合", "PV_COMBINATION_\nMODEL_DIR/meta.json", "记录管线版本、0.25分支二权重、训练截止日期和策略"],
        ["分支一树模型", "PV_V9_CONTROL_MODEL_DIR", "状态npz、清单及100个control_pv_*.joblib"],
        ["分支一时序模型", "PV_V27_MODEL_DIR", "两个随机种子参数、状态npz和neural_meta.json"],
        ["分支二", "PV_V48_MODEL_DIR", "weather.joblib、experts.joblib、state.npz、meta.json及哈希"],
        ["数据", "PV_DATA_ROOT", "定位pv_load_forecasting训练与测试目录"],
        ["实况气象", "PV_TA_WEATHER_CSV /\nTA_WEATHER_CSV", "分支二训练所需站点实况；正式推理不读取"],
        ["输出", "OUTPUT_DIR/submit_result.csv", "UTF-8 BOM；pv_id、date、pred三列"],
    ], widths=[1.05, 2.45, 3.45], font_size=9.2)
    add_body(doc, "清单文件在复用模型时逐项核对版本、必需文件、特征数、日期边界和哈希。分支二限制训练耗时的PV_V48_MAX_TRAIN_SECONDS默认值为3600秒；分支一训练线程数由PV_TRAIN_WORKERS控制，时序卷积训练设备由PV_V27_TRAIN_DEVICE指定，正式推理切换为CPU。")

    add_heading(doc, "（二）评测指标与结果分析", 2)
    add_body(doc, "官方月度考核与特殊时段考核均先逐用户计算式（1）的得分系数，再对用户取平均，分别乘20分和10分。该指标以实际值为分母，对低发电量日期的绝对误差更敏感；因此代码采用对数目标、非负截断、晴空包络、低发电风险门控和特殊时段近窗森林控制相对误差风险。")
    add_body(doc, "本地工作区未检索到该最终光伏模型的平台月度得分、极端天气专项得分或合计得分凭证；最终分支日志明确记录“ONLINE SCORE PENDING”。因此本报告不列示模型成绩对比，不把本地训练量或代码内诊断量写成官方成绩，也不作优于Baseline的判断。")
    add_body(doc, "可核验的工程结果是：顶层代码把100×31完整键集作为强约束，两个分支和融合出口均执行3100行、重复键、缺失键、意外键、有限值与非负性检查；任何一项不满足即终止输出。当前工作区未包含最终训练模型目录及对应最终提交文件，本文不把静态代码审计表述为一次实际平台推理运行。")

    add_heading(doc, "四、创新性与可行性分析", 1, page_break=False)
    add_heading(doc, "（一）技术特点", 2)
    add_body(doc, "本方案的技术特点不在于单独采用成熟的ExtraTrees、时序卷积或加权平均，而在于围绕分布式光伏业务组织可执行的特征、校正与专家协同体系。")
    add_bullet(doc, "按区域—季节—小时构建晴空辐射包络，并通过晴朗指数、辐射组成、云量和连续阴云时长描述天气状态。")
    add_bullet(doc, "以严格发布时点约束学习预报气象到实况气象的17维校正关系，同时显式输出不确定性。")
    add_bullet(doc, "联合日尺度、上午/午间/下午/夜间气象统计，连接小时天气过程与日发电量标签。")
    add_bullet(doc, "用共享专家吸收跨用户样本，用用户级专家保留装机尺度和局部响应，再以教师—学生方式对齐实况与预报信息。")
    add_bullet(doc, "融合物理特征树模型和气象校正专家模型，并对极端天气与低发电风险设置代码可追溯的校正规则。")
    add_bullet(doc, "通过模型清单、文件哈希、目录环境变量和多层键校验形成可封装的交付链。")

    add_heading(doc, "（二）工程实现与交付", 2)
    add_body(doc, "训练、推理、输出路径均由统一入口控制；历史版本源码字符串只作为嵌入实现载体，顶层实际调用函数和最终重新定义的函数决定运行逻辑。报告据此排除未被最终入口调用的HistGradientBoosting替代专家、早期模式选择和实验性版本说明。")
    add_body(doc, "方案以9个区域的气象特征共享降低重复计算，以每用户森林和局部专家保持个性，以并行树训练、批量时序推理和CPU推理满足平台运行约束。所有关键常量、模型文件和输入边界写入清单或代码检查，便于在相同数据、依赖和环境变量下复现。")

    add_heading(doc, "五、附录", 1, page_break=True)
    add_heading(doc, "（一）关键参数", 2)
    add_table_title(doc, "表5 最终调用链实际生效参数")
    add_table(doc, ["模块", "参数", "取值"], [
        ["任务", "用户/区域/预测长度", "100 / 9 / 31天"],
        ["顶层", "分支一/分支二权重", "0.75 / 0.25"],
        ["分支一森林", "树数；分裂/叶样本；最大特征；种子", "500；4/2；sqrt；42"],
        ["分支一森林", "全年/近期窗口；普通缩放", "全训练期/180天；0.965"],
        ["晴空包络", "小时分位/半窗；日分位/半窗/平滑", "0.90/10日；0.90/12样本/7点"],
        ["夜间修复", "晴空包络阈值", "5 W/m²"],
        ["历史特征", "滞后；滚动均值", "1/2/3/7/14/28；3/5/7/14/28日"],
        ["目标变换", "分支一/分支二", "ln(load+0.01×容量)；用户尺度后ln"],
        ["时序卷积", "历史窗；随机种子；普通日权重", "62天；42/2026；0.25"],
        ["风险门控", "候选天数；激活/其他倍率", "每用户最低4天；0.70/1.02"],
        ["气象桥接", "树数；叶样本；最大特征；种子", "160；3；0.8；42"],
        ["学生专家", "共享/用户树数；叶样本；组合", "300/300；6/3；0.5/0.5"],
        ["约束", "训练截止；输出", "2025-09-30；有限、非负、3100键"],
    ], widths=[1.15, 3.15, 2.65], font_size=8.9)

    add_heading(doc, "（二）关键代码片段", 2, page_break=True)
    add_body(doc, "以下片段从最终嵌入源码及顶层封装摘录，仅调整换行以适应版心，计算逻辑不变。")

    code1 = '''def fit_cs_curve(hourly_path, q=0.9, halfwin=10):
    hourly = (
        hourly_path.copy()
        if isinstance(hourly_path, pd.DataFrame)
        else pd.read_csv(hourly_path)
    )
    dt_col = (
        "valid_datetime" if "valid_datetime" in hourly.columns
        else "datetime"
    )
    hourly["datetime"] = pd.to_datetime(hourly[dt_col])
    curve = {}
    for ta, g in hourly.groupby("ta_id"):
        g = g.sort_values("datetime")
        doy = g["datetime"].dt.dayofyear.values
        hour = g["datetime"].dt.hour.values
        sr = g["SR"].astype(float).values
        arr = np.full((367, 24), np.nan)
        for i in range(len(g)):
            arr[doy[i], hour[i]] = sr[i]
        out = np.full((367, 24), np.nan)
        s = pd.DataFrame(arr[1:])
        for hh in range(24):
            col = s[hh]
            trip = pd.concat([col, col, col], ignore_index=True)
            sm = trip.rolling(
                2 * halfwin + 1, center=True, min_periods=1
            ).quantile(q)
            mid = sm.values[366:732]
            out[1:, hh] = pd.Series(mid).ffill().bfill().values
        curve[ta] = out
    return curve'''
    add_code_block(doc, "代码1 小时晴空辐射包络", code1)

    doc.add_page_break()
    code2 = '''def fit_bridge17(a, obs, dates, k):
    env = clean_fit(a[:, :k])
    raw = raw_features(a, env, dates)
    x = np.concatenate([
        raw,
        np.broadcast_to(
            np.eye(9)[:, None, :], (9, len(dates), 9)
        ),
    ], -1)
    _, mapping = map_observed(obs, fc_daily(a, dates), dates, k)
    target = transform17(np.stack([
        obs.loc[mapping[t]["station"]].reindex(dates)[FIELDS]
        .to_numpy(float) for t in TAS
    ]))
    yy = target[:, :k].reshape(-1, 17)
    xm = x[:, :k].reshape(-1, 51)
    ym = _v48_nanmean(yy, 0)
    ys = np.maximum(_v48_nanstd(yy, 0), .1)
    v = np.isfinite(yy).all(1)
    if v.sum() < 100:
        raise ValueError(
            "Insufficient complete 17-field observed weather rows"
        )
    model = ExtraTreesRegressor(
        n_estimators=160,
        min_samples_leaf=3,
        max_features=.8,
        n_jobs=4,
        random_state=42,
    ).fit(xm[v], (yy[v] - ym) / ys)
    residual = np.maximum(
        np.sqrt(np.mean(
            (model.predict(xm[v]) * ys + ym - yy[v]) ** 2, 0
        )),
        ys * .1,
    )
    return dict(
        model=model, envelope=env, mean=ym, std=ys,
        residual=residual, mapping=mapping,
        fit_end=str(dates[k - 1].date())
    ), target'''
    add_code_block(doc, "代码2 17维气象桥接模型", code2)

    doc.add_page_break()
    code3 = '''def fit_student(x, target):
    pooled = ExtraTreesRegressor(
        n_estimators=300, min_samples_leaf=6,
        max_features=.8, n_jobs=6, random_state=42
    ).fit(x, target)

    def fit_user(u):
        rows = x[:, 88] == u
        return ExtraTreesRegressor(
            n_estimators=300, min_samples_leaf=3,
            max_features=.8, n_jobs=1, random_state=42
        ).fit(x[rows], target[rows])

    local = Parallel(n_jobs=6, prefer="threads")(
        delayed(fit_user)(u) for u in range(100)
    )
    return {"pooled_tree": pooled, "local_tree": local}'''
    add_code_block(doc, "代码3 共享专家与用户级学生专家", code3)

    doc.add_page_break()
    code4 = '''def _ordered(frame):
    f = frame[["pv_id", "date", "pred"]].copy()
    f["date"] = pd.to_datetime(f.date)
    wanted = pd.MultiIndex.from_product([
        sorted("pv_" + str(i) for i in range(1, 101)),
        pd.date_range(PRED_START, PRED_END),
    ], names=["pv_id", "date"])
    if len(f) != 3100 or f.duplicated(["pv_id", "date"]).any():
        raise ValueError("PV candidate wrong row count or duplicates")
    f = f.set_index(["pv_id", "date"])
    if len(wanted.difference(f.index)) or len(f.index.difference(wanted)):
        raise ValueError("PV candidate key mismatch")
    f = f.reindex(wanted)
    if not np.isfinite(f.pred.to_numpy()).all() or (f.pred < 0).any():
        raise ValueError("PV candidate invalid prediction")
    return f

x = _ordered(a.predict_pv_model(save_output=False))
y = _ordered(b.predict_pv_model(save_output=False))
result = x.copy()
result["pred"] = (
    (1 - WEIGHT) * x.pred.to_numpy()
    + WEIGHT * y.pred.to_numpy()
)'''
    add_code_block(doc, "代码4 0.75/0.25融合与提交约束", code4)

    add_heading(doc, "（三）复现说明", 2, page_break=True)
    add_body(doc, "1. 将数据目录指向包含train_data/load_data、train_data/weather_data和test_data/weather_data的pv_load_forecasting根目录，并设置PV_DATA_ROOT。分支二训练还需通过PV_TA_WEATHER_CSV或TA_WEATHER_CSV提供ta_id站点实况天气文件。", no_indent=True)
    add_body(doc, "2. 调用submission/predict.py中的InitModel；该入口最终执行submission/utils/hyx_pv_pipeline.py的train_pv_model。若清单和全部模型文件完整且版本一致，则复用现有模型；否则按固定配方训练。", no_indent=True)
    add_body(doc, "3. 调用Detect或顶层predict_pv_model。预测阶段读取2025年10月小时预报气象，分别生成两分支结果，按预期键重排后固定融合。", no_indent=True)
    add_body(doc, "4. 输出目录由OUTPUT_DIR指定，默认写入submission/output/pv_load_forecasting/submit_result.csv。提交文件必须包含pv_id、date、pred三列，日期为YYYY/MM/DD，且为3100个唯一用户—日期键。", no_indent=True)
    add_body(doc, "5. 复现审计以各层meta.json为准：顶层清单记录组合版本，分支一记录特征数、用户映射与训练日期，分支二记录weather.joblib、experts.joblib和state.npz的SHA-256。", no_indent=True)

    add_heading(doc, "（四）参考资料", 2)
    refs = [
        "[1] 《电量预测赛道—赛题解读及评审规则说明0821》，第4、8、14—15页。",
        "[2] 《新型电力系统下电能量预测算法开发赛道—赛题数据说明文档》。",
        "[3] 《新型电力系统下电能量预测算法开发赛道—Baseline运行指南》与平台演示培训材料。",
        "[4] submission/utils/hyx_pv_pipeline.py，最终光伏训练、推理与双分支融合封装。",
        "[5] submission/predict.py，平台训练和推理入口。",
    ]
    for ref in refs:
        add_body(doc, ref, no_indent=True)

    core = doc.core_properties
    core.title = "分布式光伏发电量预测初赛技术报告"
    core.subject = "新型电力系统下电能量预测算法开发赛道—任务三"
    core.author = "参赛团队"
    core.keywords = "分布式光伏, 发电量预测, 晴空包络, 气象校正, 双分支融合"
    doc.save(DOCX_OUT)


if __name__ == "__main__":
    build_doc()
    print(DOCX_OUT)
    print(SVG_OUT)
    print(PNG_OUT)
